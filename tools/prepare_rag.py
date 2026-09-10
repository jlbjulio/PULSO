
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

import pymupdf
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
RAG_ROOT = ROOT / "data" / "rag"
MANIFEST = RAG_ROOT / "manifest.json"
OUTPUT = RAG_ROOT / "index" / "corpus.jsonl"
REPORT = RAG_ROOT / "index" / "build-report.json"
OCR_CACHE = RAG_ROOT / "index" / "ocr-cache.json"
CHARS_PER_CHUNK = 1200
OVERLAP_CHARS = 180
EMAIL = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.IGNORECASE)
INTERNATIONAL_PHONE = re.compile(r"(?<!\w)\+\d(?:[\s().-]*\d){7,14}(?!\w)")
LABELED_PHONE = re.compile(
    r"\b(?:tel(?:éfono)?|phone|fax)\s*[:.]?\s*(?:\+?\d(?:[\s().-]*\d){6,14})",
    re.IGNORECASE,
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def clean(text: str) -> str:
    redacted = EMAIL.sub("[correo omitido]", text)
    redacted = INTERNATIONAL_PHONE.sub("[teléfono omitido]", redacted)
    redacted = LABELED_PHONE.sub("[teléfono omitido]", redacted)
    return re.sub(r"\s+", " ", redacted).strip()


def pdf_pages(path: Path) -> list[tuple[str, str]]:
    document = pymupdf.open(path)
    return [
        (f"page:{index + 1}", clean(page.get_text("text"))) for index, page in enumerate(document)
    ]


def xlsx_sheets(path: Path) -> list[tuple[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[tuple[str, str]] = []
    for sheet in workbook.worksheets:
        lines = [
            " | ".join(str(value) for value in row if value is not None)
            for row in sheet.iter_rows(values_only=True)
        ]
        sections.append((f"sheet:{sheet.title}", clean("\n".join(lines))))
    return sections


def json_sections(path: Path) -> list[tuple[str, str]]:
    value = json.loads(path.read_text(encoding="utf-8"))
    return [("document", json.dumps(value, ensure_ascii=False, separators=(",", ":")))]


def chunks(text: str) -> list[str]:
    if not text:
        return []
    output: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + CHARS_PER_CHUNK, len(text))
        if end < len(text):
            boundary = text.rfind(" ", start + CHARS_PER_CHUNK // 2, end)
            if boundary > start:
                end = boundary
        output.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(start + 1, end - OVERLAP_CHARS)
    return [item for item in output if item]


def main() -> None:
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8-sig"))
    ocr_cache = json.loads(OCR_CACHE.read_text(encoding="utf-8-sig")) if OCR_CACHE.exists() else {}
    rows: list[dict[str, str]] = []
    skipped: list[dict[str, str]] = []
    source_counts: dict[str, int] = {}
    for source in manifest:
        if not source.get("include_in_rag", True):
            skipped.append({"id": source["id"], "reason": "redistribution_not_permitted"})
            continue
        path = RAG_ROOT / Path(source["local_path"])
        if not path.exists():
            skipped.append({"id": source["id"], "reason": "missing"})
            continue
        if source.get("sha256") and sha256(path) != source["sha256"]:
            skipped.append({"id": source["id"], "reason": "checksum_mismatch"})
            continue
        if path.suffix.casefold() == ".pdf":
            sections = pdf_pages(path)
        elif path.suffix.casefold() == ".xlsx":
            sections = xlsx_sheets(path)
        elif path.suffix.casefold() == ".json":
            sections = json_sections(path)
        else:
            skipped.append({"id": source["id"], "reason": "unsupported_format"})
            continue
        count = 0
        for locator, section in sections:
            if not section:
                section = ocr_cache.get(f"{source['id']}:{locator}", "")
            for index, chunk in enumerate(chunks(section), 1):
                chunk_id = f"{source['id']}:{locator}:chunk:{index}"
                header = (
                    f"[PULSO_SOURCE id={source['id']} chunk={chunk_id} locator={locator} "
                    f"title={json.dumps(source['title'], ensure_ascii=False)} "
                    f"url={source['source_url']}]"
                )
                rows.append(
                    {
                        "workspace": source["workspace"],
                        "source_id": source["id"],
                        "chunk_id": chunk_id,
                        "locator": locator,
                        "content": f"{header}\n{chunk}",
                    }
                )
                count += 1
        source_counts[source["id"]] = count
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows),
        encoding="utf-8",
    )
    report = {
        "sources": len(source_counts),
        "chunks": len(rows),
        "source_chunks": source_counts,
        "skipped": skipped,
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
