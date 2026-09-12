
from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from pulso.clinical.encounters import ClinicalDocument, Encounter
from pulso.clinical.events import ClinicalEvent, EventType
from pulso.clinical.orders import Order
from pulso.storage.audit import append_audit_event
from pulso.storage.repository import EncounterRepository

FIELD_LABELS = {
    "name": "name", "symptom": "symptom", "substance": "substance",
    "medication": "medication", "dose": "dose", "unit": "unit",
    "route": "route", "frequency": "frequency", "study": "study",
    "body_region": "region", "priority": "priority", "indication": "indication",
    "team": "team", "destination": "destination", "result": "result",
    "request": "request",
}

EVENT_LABELS = {
    EventType.PATIENT_REPORT: "Relevant history", EventType.SYMPTOM: "Symptom",
    EventType.ALLERGY: "Allergy", EventType.MEDICATION_HISTORY: "Current medication",
    EventType.VITAL_SIGN: "Vital sign", EventType.EXAM_FINDING: "Clinical finding",
    EventType.CLINICAL_ASSESSMENT: "Clinical assessment",
    EventType.DIAGNOSIS: "Documented diagnosis",
    EventType.MEDICATION_ORDER: "Medication order",
    EventType.MEDICATION_ADMINISTRATION: "Medication administered",
    EventType.PROCEDURE_ORDER: "Procedure ordered",
    EventType.PROCEDURE_PERFORMED: "Procedure performed",
    EventType.LAB_ORDER: "Laboratory order", EventType.IMAGING_ORDER: "Imaging order",
    EventType.CONSULT_ORDER: "Team or specialist request", EventType.RESULT: "Result",
    EventType.CODE_EVENT: "Critical response", EventType.TRANSFER: "Transfer",
    EventType.DISPOSITION: "Disposition", EventType.HANDOFF: "Clinical handoff",
}

ORDER_STATE_LABELS = {
    "awaiting_confirmation": "Awaiting confirmation", "confirmed": "Confirmed",
    "dispatched": "Dispatched", "accepted": "Accepted", "in_progress": "In progress",
    "completed": "Completed", "cancelled": "Cancelled", "failed": "Disconnected",
}

COLORS = {
    "navy": "102A25",
    "green": "1B7059",
    "mint": "DDF4EB",
    "pale": "F3F8F6",
    "line": "CADBD5",
    "ink": "16332D",
    "muted": "60766F",
    "white": "FFFFFF",
    "amber": "FFF3DF",
}


def _label(event: ClinicalEvent) -> str:
    payload = ", ".join(
        f"{FIELD_LABELS.get(key, key.replace('_', ' '))}: {value}"
        for key, value in event.payload.items()
        if key != "explicit_command"
    )
    return payload or event.type.value


def build_document(encounter_id: str, events: list[ClinicalEvent]) -> ClinicalDocument:
    groups: dict[str, list[str]] = {
        "allergies": [],
        "symptoms": [],
        "findings": [],
        "assessments": [],
        "interventions": [],
        "results": [],
        "pending_orders": [],
    }
    for event in events:
        value = _label(event)
        if event.type == EventType.ALLERGY:
            groups["allergies"].append(value)
        elif event.type in {EventType.SYMPTOM, EventType.PATIENT_REPORT}:
            groups["symptoms"].append(value)
        elif event.type in {EventType.VITAL_SIGN, EventType.EXAM_FINDING}:
            groups["findings"].append(value)
        elif event.type in {EventType.CLINICAL_ASSESSMENT, EventType.DIAGNOSIS}:
            groups["assessments"].append(value)
        elif event.type in {
            EventType.MEDICATION_ADMINISTRATION,
            EventType.PROCEDURE_PERFORMED,
            EventType.CODE_EVENT,
        }:
            groups["interventions"].append(value)
        elif event.type == EventType.RESULT:
            groups["results"].append(value)
        elif event.actionable:
            groups["pending_orders"].append(value)
    summary_parts = [
        f"Symptoms: {'; '.join(groups['symptoms']) or 'no data'}.",
        f"Findings: {'; '.join(groups['findings']) or 'no data'}.",
        f"Documented assessment: {'; '.join(groups['assessments']) or 'no data'}.",
    ]
    return ClinicalDocument(encounter_id=encounter_id, summary=" ".join(summary_parts), **groups)


def save_document(
    repository: EncounterRepository,
    document: ClinicalDocument,
    *,
    actor: str,
    signature: str | None = None,
) -> ClinicalDocument:
    now = datetime.now(UTC)
    if signature:
        document = document.model_copy(update={"signed_by": actor, "signed_at": now})
    payload = document.model_dump(mode="json")
    with repository.database.transaction() as connection:
        connection.execute(
            """INSERT INTO clinical_documents
            (encounter_id, document_json, signed_by, signed_at, updated_at)
            VALUES (?, ?, ?, ?, ?)
            ON CONFLICT(encounter_id) DO UPDATE SET document_json=excluded.document_json,
            signed_by=excluded.signed_by, signed_at=excluded.signed_at,
            updated_at=excluded.updated_at""",
            (
                document.encounter_id,
                json.dumps(payload, ensure_ascii=False),
                document.signed_by,
                document.signed_at.isoformat() if document.signed_at else None,
                now.isoformat(),
            ),
        )
        append_audit_event(
            connection,
            entity_type="clinical_document",
            entity_id=document.encounter_id,
            action="signed" if signature else "reviewed",
            actor=actor,
            payload=payload,
        )
    return document


def _shade(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _cell_margins(
    cell,
    top: int = 110,
    start: int = 130,
    bottom: int = 110,
    end: int = 130,
) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_text(
    cell,
    text: str,
    *,
    size: float = 9,
    color: str = COLORS["ink"],
    bold: bool = False,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _cell_margins(cell)


def _section_heading(report: Document, title: str) -> None:
    paragraph = report.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(13)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(title.upper())
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(COLORS["green"])


def _information_card(report: Document, values: list[str]) -> None:
    table = report.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    _shade(cell, COLORS["pale"])
    _cell_margins(cell, 130, 160, 130, 160)
    cell.text = ""
    for index, value in enumerate(values or ["No information documented."]):
        paragraph = cell.add_paragraph() if index else cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(4 if index < len(values) - 1 else 0)
        bullet = paragraph.add_run("●  ")
        bullet.font.color.rgb = RGBColor.from_string(COLORS["green"])
        bullet.font.size = Pt(6)
        text = paragraph.add_run(value)
        text.font.name = "Aptos"
        text.font.size = Pt(9)
        text.font.color.rgb = RGBColor.from_string(COLORS["ink"])


def _page_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = field
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def export_word_report(
    encounter: Encounter,
    document: ClinicalDocument,
    events: list[ClinicalEvent],
    orders: list[Order],
    output: str | Path,
) -> Path:
    destination = Path(output).resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    report = Document()
    report.core_properties.title = "Emergency encounter report"
    report.core_properties.subject = encounter.id
    report.core_properties.author = "PULSO"
    section = report.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.22)
    section.footer_distance = Inches(0.28)
    normal = report.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor.from_string(COLORS["ink"])
    normal.paragraph_format.space_after = Pt(5)

    header = section.header
    header_table = header.add_table(rows=1, cols=2, width=Inches(7.1))
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.columns[0].width = Inches(4.9)
    header_table.columns[1].width = Inches(2.2)
    _shade(header_table.cell(0, 0), COLORS["navy"])
    _shade(header_table.cell(0, 1), COLORS["navy"])
    _set_cell_text(
        header_table.cell(0, 0),
        "PULSO",
        size=14,
        color=COLORS["white"],
        bold=True,
    )
    _set_cell_text(
        header_table.cell(0, 1),
        "EMERGENCY DEPARTMENT",
        size=8,
        color="A7D9C9",
        bold=True,
    )
    header_table.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    eyebrow = report.add_paragraph()
    eyebrow.paragraph_format.space_before = Pt(7)
    eyebrow.paragraph_format.space_after = Pt(2)
    eyebrow_run = eyebrow.add_run("ASSISTED CLINICAL RECORD")
    eyebrow_run.bold = True
    eyebrow_run.font.size = Pt(8)
    eyebrow_run.font.color.rgb = RGBColor.from_string(COLORS["green"])
    title = report.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    title_run = title.add_run("Emergency encounter report")
    title_run.bold = True
    title_run.font.name = "Aptos Display"
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor.from_string(COLORS["navy"])
    subtitle = report.add_paragraph(
        f"Encounter {encounter.id}  ·  Closed "
        f"{datetime.now().astimezone().strftime('%d/%m/%Y %H:%M')}"
    )
    subtitle.paragraph_format.space_after = Pt(10)
    subtitle.runs[0].font.size = Pt(8)
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(COLORS["muted"])

    metadata = report.add_table(rows=2, cols=4)
    metadata.alignment = WD_TABLE_ALIGNMENT.CENTER
    metadata.autofit = False
    metadata_values = [
        ("PATIENT", encounter.patient_ref), ("TREATMENT BAY", encounter.bed),
        ("CLINICIAN", encounter.clinician_id),
        ("STARTED", encounter.started_at.astimezone().strftime("%d/%m/%Y %H:%M")),
    ]
    for index, (label, value) in enumerate(metadata_values):
        label_cell = metadata.cell(0, index)
        value_cell = metadata.cell(1, index)
        _shade(label_cell, COLORS["mint"])
        _shade(value_cell, COLORS["pale"])
        _set_cell_text(label_cell, label, size=7, color=COLORS["green"], bold=True)
        _set_cell_text(value_cell, value, size=9, color=COLORS["ink"], bold=True)

    _section_heading(report, "Clinical summary")
    summary_table = report.add_table(rows=1, cols=1)
    summary_cell = summary_table.cell(0, 0)
    _shade(summary_cell, COLORS["mint"])
    _cell_margins(summary_cell, 150, 170, 150, 170)
    _set_cell_text(summary_cell, document.summary, size=10, color=COLORS["ink"])

    sections = [
        ("Allergies", document.allergies), ("Symptoms and relevant history", document.symptoms),
        ("Vital signs and findings", document.findings),
        ("Assessments and diagnoses", document.assessments),
        ("Completed interventions", document.interventions), ("Results", document.results),
    ]
    for heading, values in sections:
        _section_heading(report, heading)
        _information_card(report, values)

    _section_heading(report, "Orders and coordination")
    if orders:
        order_table = report.add_table(rows=1, cols=3)
        order_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        order_table.autofit = False
        widths = [Inches(1.55), Inches(4.15), Inches(1.4)]
        for cell, value, width in zip(
            order_table.rows[0].cells,
            ["DESTINATION", "REQUEST", "STATUS"],
            widths,
            strict=True,
        ):
            cell.width = width
            _shade(cell, COLORS["navy"])
            _set_cell_text(cell, value, size=7, color=COLORS["white"], bold=True)
        for index, order in enumerate(orders):
            cells = order_table.add_row().cells
            background = COLORS["white"] if index % 2 == 0 else COLORS["pale"]
            values = [
                order.destination,
                order.request,
                ORDER_STATE_LABELS.get(order.state.value, order.state.value),
            ]
            for cell, value, width in zip(cells, values, widths, strict=True):
                cell.width = width
                _shade(cell, background)
                _set_cell_text(cell, value, size=8.3)
    else:
        _information_card(report, ["No orders were recorded during this encounter."])

    _section_heading(report, "Clinical timeline")
    if events:
        timeline = report.add_table(rows=1, cols=3)
        timeline.alignment = WD_TABLE_ALIGNMENT.CENTER
        timeline.autofit = False
        timeline_widths = [Inches(0.7), Inches(1.9), Inches(4.5)]
        for cell, value, width in zip(
            timeline.rows[0].cells,
            ["TIME", "EVENT", "DETAIL"],
            timeline_widths,
            strict=True,
        ):
            cell.width = width
            _shade(cell, COLORS["green"])
            _set_cell_text(cell, value, size=7, color=COLORS["white"], bold=True)
        for index, event in enumerate(events):
            cells = timeline.add_row().cells
            background = COLORS["white"] if index % 2 == 0 else COLORS["pale"]
            values = [
                event.created_at.astimezone().strftime("%H:%M:%S"),
                EVENT_LABELS.get(event.type, event.type.value),
                _label(event),
            ]
            for cell, value, width in zip(cells, values, timeline_widths, strict=True):
                cell.width = width
                _shade(cell, background)
                _set_cell_text(cell, value, size=8)
    else:
        _information_card(report, ["No relevant clinical events were recorded."])

    _section_heading(report, "Clinician validation")
    signature_table = report.add_table(rows=1, cols=2)
    signature_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    signature_values = [
        ("REVIEWED AND SIGNED BY", document.signed_by or "Pending"),
        (
            "VALIDATION DATE",
            document.signed_at.astimezone().strftime("%d/%m/%Y %H:%M")
            if document.signed_at
            else "Pending",
        ),
    ]
    for cell, (label, value) in zip(signature_table.rows[0].cells, signature_values, strict=True):
        _shade(cell, COLORS["amber"])
        _cell_margins(cell, 140, 160, 140, 160)
        cell.text = ""
        label_paragraph = cell.paragraphs[0]
        label_run = label_paragraph.add_run(label)
        label_run.bold = True
        label_run.font.size = Pt(7)
        label_run.font.color.rgb = RGBColor.from_string(COLORS["muted"])
        value_paragraph = cell.add_paragraph()
        value_paragraph.paragraph_format.space_after = Pt(0)
        value_run = value_paragraph.add_run(value)
        value_run.bold = True
        value_run.font.size = Pt(10)
        value_run.font.color.rgb = RGBColor.from_string(COLORS["ink"])

    notice = report.add_paragraph()
    notice.paragraph_format.space_before = Pt(9)
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice_run = notice.add_run(
        "Operational support document. Requires clinical validation "
        "and does not replace the official health record."
    )
    notice_run.italic = True
    notice_run.font.size = Pt(7.5)
    notice_run.font.color.rgb = RGBColor.from_string(COLORS["muted"])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("PULSO  ·  CONFIDENTIAL  ·  Page ")
    footer_run.font.size = Pt(7)
    footer_run.font.color.rgb = RGBColor.from_string(COLORS["muted"])
    _page_field(footer, "PAGE")
    footer.add_run(" of ")
    _page_field(footer, "NUMPAGES")
    report.save(destination)
    return destination
